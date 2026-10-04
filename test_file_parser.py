import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter

from backend.exceptions import DocumentParseError
from backend.parsers.file_parser import FileParser
from file_reader import extract_text_from_upload


class FileParserTests(unittest.TestCase):
    def setUp(self):
        self.parser = FileParser()

    def assert_error(self, data, filename, code):
        with self.assertRaises(DocumentParseError) as caught:
            self.parser.parse(data, filename)
        self.assertEqual(caught.exception.code, code)

    def test_txt_and_upload(self):
        upload = SimpleNamespace(name="cv.TXT", getvalue=lambda: b"\xef\xbb\xbfSkills\nPython")
        result = self.parser.parse(upload)
        self.assertEqual(result.raw_text, "Skills\nPython")
        self.assertEqual(result.source_type, "txt")
        self.assertEqual(result.filename, "cv.TXT")

    def test_latin1_txt(self):
        self.assertIn("Andr\u00e9", self.parser.parse("Andr\u00e9 CV".encode("latin-1"), "cv.txt").raw_text)

    def test_empty_and_whitespace(self):
        self.assert_error(b"", "cv.txt", "EMPTY_DOCUMENT")
        self.assert_error(b" \n\t", "cv.txt", "EMPTY_DOCUMENT")

    def test_unsupported(self):
        self.assert_error(b"CV", "cv.jpg", "UNSUPPORTED_FILE")

    def test_read_error_is_safe(self):
        with self.assertRaises(DocumentParseError) as caught:
            self.parser.parse("missing-cv.txt")
        self.assertEqual(caught.exception.code, "FILE_READ_FAILED")

    def test_invalid_pdf_and_docx(self):
        self.assert_error(b"not a PDF", "cv.pdf", "PDF_TEXT_EXTRACTION_FAILED")
        self.assert_error(b"not a DOCX", "cv.docx", "DOCX_TEXT_EXTRACTION_FAILED")

    def test_scanned_pdf_empty(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        data = BytesIO()
        writer.write(data)
        self.assert_error(data.getvalue(), "cv.pdf", "EMPTY_DOCUMENT")

    def test_low_extraction_quality(self):
        pages = [SimpleNamespace(extract_text=lambda: "a" * 40) for _ in range(2)]
        with patch("pypdf.PdfReader", return_value=SimpleNamespace(pages=pages)):
            result = self.parser.parse(b"%PDF-mock", "cv.pdf")
        self.assertEqual(result.page_count, 2)
        self.assertEqual(result.extraction_quality, "low")

    def test_good_extraction_quality(self):
        pages = [SimpleNamespace(extract_text=lambda: "CV words " * 100)]
        with patch("pypdf.PdfReader", return_value=SimpleNamespace(pages=pages)):
            self.assertEqual(self.parser.parse(b"%PDF-mock", "cv.pdf").extraction_quality, "good")

    def test_docx_paragraph_and_table_order(self):
        doc = Document()
        doc.add_paragraph("Education")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Bachelor"
        table.cell(0, 1).text = "Computer Science"
        doc.add_paragraph("Skills: Python")
        data = BytesIO()
        doc.save(data)
        result = self.parser.parse(data.getvalue(), "cv.docx")
        self.assertLess(result.raw_text.index("Bachelor"), result.raw_text.index("Skills"))
        self.assertIn("Computer Science", result.raw_text)

    def test_legacy_upload_adapter(self):
        upload = SimpleNamespace(name="cv.txt", getvalue=lambda: b"Skills: Python")
        self.assertEqual(extract_text_from_upload(upload), ("Skills: Python", None))
        self.assertEqual(extract_text_from_upload(None), ("", None))
        _, error = extract_text_from_upload(SimpleNamespace(name="cv.jpg", getvalue=lambda: b"private content"))
        self.assertIn("UNSUPPORTED_FILE", error)
        self.assertNotIn("private content", error)
