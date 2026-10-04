from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pydantic import ValidationError
import pytest
from streamlit.testing.v1 import AppTest

from backend.exceptions import InvalidCVError
from backend.parsers.cv_parser import extract_sections
from backend.services.application_documents import (
    LetterRequest, ResumeDraft, cover_letter, evidence_options, letter_docx, resume_docx, resume_text,
)
from tests.helpers import CV, JD, PYTHON_QUOTE, SQL_QUOTE

APP = str(Path(__file__).resolve().parents[1] / "app.py")


def draft(**changes):
    return ResumeDraft(full_name="Nguyễn Minh An", email="an@example.test", headline="Data Analyst",
                       **{"language": "vi", "projects": "- Xây dựng báo cáo bằng Python.", **changes})


def request(**changes):
    return LetterRequest(**{"full_name": "Nguyễn Minh An", "company": "Example Co", "role": "Data Analyst",
                            "cv_text": CV, "jd_text": JD, "evidence": [PYTHON_QUOTE], **changes})


@pytest.mark.parametrize("language,heading", [("vi", "Dự án"), ("en", "Projects")])
def test_resume_text_and_word_keep_facts_and_parseable_headings(language, heading):
    model = draft(language=language)
    text = resume_text(model)
    word = Document(BytesIO(resume_docx(model)))
    paragraphs = [p.text for p in word.paragraphs]
    assert paragraphs[0] == "Nguyễn Minh An"
    assert heading in paragraphs
    assert "Xây dựng báo cáo bằng Python." in paragraphs
    assert "an@example.test" in paragraphs
    assert "Python" in extract_sections(text)["projects"]
    assert "Education" not in paragraphs and "Học vấn" not in paragraphs
    assert not word.tables
    assert not any(p.text for p in word.sections[0].header.paragraphs)


def test_templates_differ_without_losing_content():
    modern = Document(BytesIO(resume_docx(draft(template="modern"))))
    classic = Document(BytesIO(resume_docx(draft(template="classic"))))
    assert [p.text for p in modern.paragraphs] == [p.text for p in classic.paragraphs]
    assert classic.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert modern.styles["Normal"].font.name != classic.styles["Normal"].font.name
    assert not modern.styles.element.xpath(".//w:pBdr")
    assert not classic.styles["Title"].element.xpath(".//w:rFonts/@w:asciiTheme")


@pytest.mark.parametrize("data", [{"full_name": " ", "skills": "Python"},
    {"full_name": "An"}, {"full_name": "An", "skills": "Python", "email": "bad@"},
    {"full_name": "An", "skills": "Python\x00"}])
def test_invalid_resume_is_rejected(data):
    with pytest.raises(ValidationError):
        ResumeDraft(**data)


def test_evidence_ranking_ignores_benefits_and_excludes_contact():
    options = evidence_options(CV, "Requirements\nPython required\nBenefits\nSQL training provided")
    assert PYTHON_QUOTE in options[:2]
    assert options.index(PYTHON_QUOTE) < options.index(SQL_QUOTE)
    assert not any("@" in option for option in options)


@pytest.mark.parametrize("quotes", [["Managed 50 people."], [PYTHON_QUOTE, PYTHON_QUOTE], [],
    [PYTHON_QUOTE, SQL_QUOTE, "Made up result"]])
def test_letter_rejects_unverified_or_duplicate_evidence(quotes):
    with pytest.raises(ValidationError):
        request(evidence=quotes)


@pytest.mark.parametrize("language", ["en", "vi"])
def test_letter_keeps_exact_evidence_role_and_company(language):
    text = cover_letter(request(language=language, motivation="I enjoy building useful data tools."))
    assert PYTHON_QUOTE in text
    assert "Example Co" in text and "Data Analyst" in text and "Nguyễn Minh An" in text
    assert "I enjoy building useful data tools." in text
    assert "2 years" not in text  # A JD requirement must not become an applicant claim.
    assert "20%" not in text
    edited = text + "\n\nAvailable for a conversation next week."
    paragraphs = [p.text for p in Document(BytesIO(letter_docx(edited))).paragraphs]
    assert paragraphs[-1] == "Available for a conversation next week."


@pytest.mark.parametrize("text", [" ", "bad\x00text", "a" * 20001])
def test_letter_export_validates_edited_text(text):
    with pytest.raises(InvalidCVError):
        letter_docx(text)


def app_at(page):
    app = AppTest.from_file(APP)
    app.session_state["page"] = page
    return app.run(timeout=20)


def click(app, label):
    next(button for button in app.button if button.label == label).click().run(timeout=20)
    assert not app.exception


def test_builder_submit_navigation_and_analyzer_handoff():
    with patch("frontend.review_ui.AnalyzerReviewService") as analyzer:
        app = app_at("CV Builder")
        app.text_input(key="builder_full_name").set_value("Nguyễn Minh An")
        app.text_input(key="builder_headline").set_value("Data Analyst")
        app.text_area(key="builder_projects").set_value("- Built a Python dashboard.")
        click(app, "Build CV")
        text = resume_text(ResumeDraft.model_validate(app.session_state["builder_submitted"]))
        assert "Python dashboard" in text
        click(app, "Dashboard")
        click(app, "CV Builder")
        assert app.text_input(key="builder_full_name").value == "Nguyễn Minh An"
        app.session_state["saved_cv_id"] = "old-cv"
        app.session_state["review_comparison"] = {"old": True}
        app.session_state["jd_text"] = JD
        click(app, "Use this CV in Analyzer")
        assert app.session_state["page"] == "Analyzer"
        assert app.session_state["cv_document"].raw_text == text
        assert app.text_area[0].value == text
        assert app.text_area[1].value == JD
        assert app.session_state["target_position"] == "Data Analyst"
        assert "saved_cv_id" not in app.session_state
        assert "review_comparison" not in app.session_state
        analyzer.assert_not_called()


def test_cover_letter_edits_survive_navigation_and_stale_target_blocks_download():
    app = app_at("Cover Letter")
    app.text_area(key="letter_cv_text").set_value(CV)
    app.text_area(key="letter_jd_text").set_value(JD)
    app.text_input(key="letter_full_name").set_value("Student Name")
    app.text_input(key="letter_role").set_value("Data Analyst")
    app.text_input(key="letter_company").set_value("Example Co")
    click(app, "Create letter draft")
    assert len(app.get("download_button")) == 2
    edited = app.text_area(key="letter_editor").value + "\n\nMy own additional sentence."
    app.text_area(key="letter_editor").set_value(edited).run()
    click(app, "Dashboard")
    click(app, "Cover Letter")
    assert app.text_area(key="letter_editor").value == edited
    app.text_input(key="letter_company").set_value("Different Company").run()
    assert not app.exception
    assert len(app.get("download_button")) == 0
    assert any("Inputs changed" in item.value for item in app.warning)
    click(app, "Create letter draft")
    assert "Different Company" in app.text_area(key="letter_editor").value
    assert "Example Co" not in app.text_area(key="letter_editor").value
    assert len(app.get("download_button")) == 2


def test_changing_cv_removes_stale_evidence():
    app = app_at("Cover Letter")
    app.text_area(key="letter_cv_text").set_value(CV).run()
    assert app.multiselect[0].value
    app.text_area(key="letter_cv_text").set_value("Projects\n- Designed a new web application with React.").run()
    assert not app.exception
    assert app.multiselect[0].value == ["Designed a new web application with React."]
